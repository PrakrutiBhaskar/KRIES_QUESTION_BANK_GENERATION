"""
Paper -> Word (.docx) bytes.

Mirrors the layout of the PDF renderers (header, instructions, optional
sections, questions with marks on the right, options / match table, figures,
then the Answer Key on a fresh page) so the two formats read the same, but the
result stays fully editable in Word.

Kannada: runs carry both the Latin and the complex-script font names, so Word
shapes Kannada text correctly (it uses the complex-script font slot for it).
"""
from __future__ import annotations

import logging
import re
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from ..figures import answer_key_figure, fit_size_mm, loaded_figure, read_figure_bytes
from .answer_format import (
    MATCH_COLUMN_HEADINGS,
    choice_note,
    format_answer,
    match_key_label,
    match_layout,
    split_label,
)
from .html import item_section, ordered_items, section_totals

logger = logging.getLogger("backend.export")

_OPTION_LABELS = "abcdefgh"
_BODY_FONT = "Noto Sans"
_KANNADA_FONT = "Noto Sans Kannada"
_GREY = RGBColor(0x44, 0x44, 0x44)
_BLUE = RGBColor(0x2B, 0x3A, 0x67)

# Characters XML 1.0 cannot hold would make python-docx raise.
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(text) -> str:
    return _ILLEGAL_XML.sub("", str(text if text is not None else ""))


def _instruction_lines(types: set[str]) -> list[str]:
    lines = ["All questions are compulsory.", "Marks for each question are shown against it."]
    if "MCQ" in types:
        lines.append("For multiple-choice questions, write the letter of the correct option.")
    if "Long" in types:
        lines.append("Answer long-answer questions in full, showing all steps or points.")
    if "Fill" in types:
        lines.append("For fill-in-the-blank questions, write the missing word or phrase.")
    if "Match" in types:
        lines.append(
            "For match-the-following questions, write the letter of the matching "
            "option against each number."
        )
    return lines


def _set_fonts(run, *, size: float | None = None, bold=None, italic=None, color=None) -> None:
    run.font.name = _BODY_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), _BODY_FONT)
    rfonts.set(qn("w:hAnsi"), _BODY_FONT)
    rfonts.set(qn("w:cs"), _KANNADA_FONT)
    rfonts.set(qn("w:eastAsia"), _BODY_FONT)
    if size is not None:
        run.font.size = Pt(size)
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(int(size * 2)))
        rpr.append(szcs)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _para(doc_or_cell, text="", *, size=11, bold=False, italic=False, color=None,
          align=None, left_mm=0.0, space_after=3, keep_next=False):
    p = doc_or_cell.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.space_before = Pt(0)
    if left_mm:
        fmt.left_indent = Mm(left_mm)
    if align is not None:
        p.alignment = align
    if keep_next:
        fmt.keep_with_next = True
    if text:
        _set_fonts(p.add_run(_clean(text)), size=size, bold=bold, italic=italic, color=color)
    return p


def _bottom_border(paragraph, size=8) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "444444")
    borders.append(bottom)
    ppr.append(borders)


def _add_figure(doc, figure, *, left_mm: float = 0.0) -> None:
    """Insert a figure (and caption); a missing/unreadable image is skipped."""
    data = read_figure_bytes(figure)
    if data is None:
        return
    w, h = fit_size_mm(figure.width, figure.height)
    p = _para(doc, left_mm=left_mm, space_after=2, keep_next=bool(figure.caption))
    try:
        p.add_run().add_picture(BytesIO(data), width=Mm(w), height=Mm(h))
    except Exception as exc:  # unsupported/corrupt image must not fail the export
        logger.warning("Could not embed figure %s in the Word export: %s", figure.id, exc)
        p._p.getparent().remove(p._p)
        return
    if figure.caption:
        _para(doc, figure.caption, size=9, italic=True, color=_GREY, left_mm=left_mm)


def _page_number_footer(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def field(code: str):
        run = p.add_run()
        _set_fonts(run, size=9, color=RGBColor(0x66, 0x66, 0x66))
        for kind, text in (("begin", None), (None, code), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)

    _set_fonts(p.add_run("Page "), size=9, color=RGBColor(0x66, 0x66, 0x66))
    field(" PAGE ")
    _set_fonts(p.add_run(" of "), size=9, color=RGBColor(0x66, 0x66, 0x66))
    field(" NUMPAGES ")


def _shade(cell, hex_fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _question_line(doc, n: int, text: str, marks: int) -> None:
    """'1.  Question text ............ [2]' with the marks on a right-aligned tab."""
    p = _para(doc, space_after=3, keep_next=True)
    p.paragraph_format.left_indent = Mm(9)
    p.paragraph_format.first_line_indent = Mm(-9)
    usable_mm = 210 - 16 - 16
    p.paragraph_format.tab_stops.add_tab_stop(Mm(9), WD_TAB_ALIGNMENT.LEFT)
    p.paragraph_format.tab_stops.add_tab_stop(Mm(usable_mm), WD_TAB_ALIGNMENT.RIGHT)
    _set_fonts(p.add_run(f"{n}.\t"), size=11, bold=True)
    _set_fonts(p.add_run(_clean(text)), size=11)
    _set_fonts(p.add_run(f"\t[{marks}]"), size=11, bold=True)


def _match_table(doc, rows) -> None:
    h1, h2 = MATCH_COLUMN_HEADINGS
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, head in zip(table.rows[0].cells, (h1, h2)):
        cell.text = ""
        _set_fonts(cell.paragraphs[0].add_run(head), size=10.5, bold=True)
        _shade(cell, "EEF0F6")
    for a, b in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, (a, b)):
            cell.text = ""
            _set_fonts(cell.paragraphs[0].add_run(_clean(value)), size=10.5)
    _para(doc, space_after=4)


def render_docx(paper, *, include_answer_key: bool = True) -> bytes:
    """Build the paper as a Word document and return its bytes."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(16)
    section.top_margin, section.bottom_margin = Mm(18), Mm(20)
    _page_number_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = _BODY_FONT
    normal.font.size = Pt(11)

    grades = sorted({item.question.grade for item in paper.items})
    grade_label = ", ".join(str(g) for g in grades) if grades else "-"

    # --- header --------------------------------------------------------
    _para(doc, "Karnataka State Board", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(doc, paper.title, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    meta = _para(doc, space_after=4)
    meta.paragraph_format.tab_stops.add_tab_stop(Mm(89), WD_TAB_ALIGNMENT.CENTER)
    meta.paragraph_format.tab_stops.add_tab_stop(Mm(178), WD_TAB_ALIGNMENT.RIGHT)
    _set_fonts(
        meta.add_run(
            f"Subject: {_clean(paper.subject.name)}\tGrade: {grade_label}\t"
            f"Maximum Marks: {paper.total_marks}"
        ),
        size=10.5,
        bold=True,
    )
    _bottom_border(meta)

    types = {item.question.type.value for item in paper.items}
    _para(doc, "Instructions", size=10.5, bold=True, space_after=1)
    for line in _instruction_lines(types):
        _para(doc, line, size=10, color=_GREY, left_mm=3, space_after=0)
    _para(doc, space_after=4)

    ordered = ordered_items(paper)
    section_marks = section_totals(ordered)

    # --- questions -----------------------------------------------------
    current_section = None
    for n, item in enumerate(ordered, start=1):
        q = item.question
        sec = item_section(item)
        if sec and sec != current_section:
            current_section = sec
            head = _para(doc, space_after=4, keep_next=True)
            head.paragraph_format.space_before = Pt(6)
            head.paragraph_format.tab_stops.add_tab_stop(Mm(178), WD_TAB_ALIGNMENT.RIGHT)
            _set_fonts(head.add_run(f"{_clean(sec)}\t{section_marks[sec]} marks"), size=11.5, bold=True, color=_BLUE)
            _bottom_border(head, size=6)

        match = match_layout(q.text, q.options) if q.type.value == "Match" else None
        _question_line(doc, n, match[0] if match else q.text, item.effective_marks)
        _add_figure(doc, loaded_figure(q, "figure"), left_mm=9)
        if match:
            _match_table(doc, match[1])
        elif q.options:
            for label, option in zip(_OPTION_LABELS, q.options):
                _para(doc, f"({label}) {option}", size=10.5, left_mm=12, space_after=1)
            _para(doc, space_after=3)
        else:
            _para(doc, space_after=4)

    # --- answer key ----------------------------------------------------
    if include_answer_key:
        doc.add_page_break()
        _para(doc, "Answer Key", size=14, bold=True, space_after=4)
        for n, item in enumerate(ordered, start=1):
            q = item.question
            a_fig = answer_key_figure(q)
            if q.type.value == "MCQ":
                _para(doc, f"{n}. {q.answer}", size=11, bold=True, space_after=1, keep_next=bool(q.explanation))
                if q.explanation:
                    _para(doc, q.explanation, size=10, color=_GREY, left_mm=6, space_after=3)
            elif q.type.value == "Match":
                _para(doc, f"{n}. {match_key_label(q.answer)}", size=11, bold=True, space_after=3)
            else:
                fa = format_answer(
                    q.answer, item.effective_marks, q.type.value, has_figure=a_fig is not None
                )
                if len(fa.points) > 1:
                    note = choice_note(len(fa.points), item.effective_marks)
                    p = _para(doc, space_after=1, keep_next=True)
                    _set_fonts(p.add_run(f"{n}. "), size=11, bold=True)
                    if fa.lead:
                        _set_fonts(p.add_run(_clean(fa.lead) + " "), size=11, bold=True)
                    elif not note:
                        _set_fonts(p.add_run(f"[{item.effective_marks} marks] "), size=11, bold=True)
                    if note:
                        # e.g. "[2 marks – any 2 of the following 5 points]"
                        _set_fonts(p.add_run(_clean(note)), size=10, bold=not fa.lead, italic=bool(fa.lead), color=_BLUE if fa.lead else None)
                    for i, point in enumerate(fa.points, start=1):
                        pp = _para(doc, left_mm=14, space_after=1)
                        pp.paragraph_format.first_line_indent = Mm(-6)
                        _set_fonts(pp.add_run(f"{i}.  {_clean(point)}"), size=10.5)
                else:
                    text = fa.points[0] if fa.points else q.answer
                    _para(doc, f"{n}. {text}", size=11, space_after=2)
                if fa.reference:
                    _para(doc, fa.reference, size=9.5, italic=True, color=_GREY, left_mm=6, space_after=1)
                if fa.split:
                    _para(doc, split_label(fa.split), size=9.5, bold=True, color=_BLUE, left_mm=6, space_after=2)
            _add_figure(doc, a_fig, left_mm=6)
            _para(doc, space_after=3)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
