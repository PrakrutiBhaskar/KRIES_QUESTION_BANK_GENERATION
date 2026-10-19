"""Hindi (Devanagari) PDF/DOCX export must shape conjuncts and not need WeasyPrint."""
from io import BytesIO
from types import SimpleNamespace as NS

import pytest

from app.config import settings
from app.services.export import renderer
from app.services.export.fpdf_renderer import fpdf_available

pytestmark = pytest.mark.skipif(not fpdf_available(), reason="fpdf2/uharfbuzz/fonts not available")

HINDI_Q = "\u0926\u093f\u0932\u094d\u0932\u0940 \u0915\u093f\u0938 \u0926\u0947\u0936 \u0915\u0940 \u0930\u093e\u091c\u0927\u093e\u0928\u0940 \u0939\u0948?"
HINDI_A = "1. \u0926\u093f\u0932\u094d\u0932\u0940 \u092d\u093e\u0930\u0924 \u0915\u0940 \u0930\u093e\u091c\u0927\u093e\u0928\u0940 \u0939\u0948\u0964 2. \u092f\u0939\u093e\u0901 \u0932\u093e\u0932 \u0915\u093f\u0932\u093e \u0939\u0948\u0964 3. \u092f\u0939 \u090f\u0915 \u0910\u0924\u093f\u0939\u093e\u0938\u093f\u0915 \u0938\u094d\u0925\u093e\u0928 \u0939\u0948\u0964"


def _paper(subject="Hindi", text=HINDI_Q):
    q = NS(
        type=NS(value="Short"), text=text, options=None, grade=7,
        answer=HINDI_A, explanation="",
    )
    item = NS(question=q, order_index=0, effective_marks=3)
    return NS(title="\u0939\u093f\u0902\u0926\u0940 \u092a\u0930\u0940\u0915\u094d\u0937\u093e", subject=NS(name=subject), items=[item], total_marks=3)


@pytest.mark.parametrize("choice", ["auto", "reportlab", "fpdf"])
def test_hindi_paper_renders_instead_of_erroring(monkeypatch, choice):
    monkeypatch.setattr(settings, "pdf_renderer", choice)
    monkeypatch.setattr(renderer, "weasyprint_available", lambda: False)
    pdf = renderer.render_pdf(_paper())
    assert pdf.startswith(b"%PDF") and len(pdf) > 5000


def test_hindi_text_in_another_subject_is_routed_to_fpdf(monkeypatch):
    monkeypatch.setattr(settings, "pdf_renderer", "reportlab")
    monkeypatch.setattr(renderer, "weasyprint_available", lambda: False)
    pdf = renderer.render_pdf(_paper(subject="Social Science"))
    assert pdf.startswith(b"%PDF")


def test_reportlab_refuses_devanagari_directly():
    with pytest.raises(renderer.ServiceUnavailableError):
        renderer._render_reportlab(_paper())


def test_docx_uses_devanagari_font_for_complex_script_runs():
    from docx import Document
    from docx.oxml.ns import qn

    from app.services.export import docx_renderer

    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run(HINDI_Q)
    docx_renderer._set_fonts(run, size=11)
    latin = p.add_run("What is the capital?")
    docx_renderer._set_fonts(latin, size=11)
    cs = lambda r: r._element.rPr.find(qn("w:rFonts")).get(qn("w:cs"))
    assert cs(run) == "Noto Sans Devanagari"
    assert cs(latin) == "Noto Sans Kannada"
