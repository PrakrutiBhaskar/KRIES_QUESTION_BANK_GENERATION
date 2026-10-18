"""Word (.docx) export: same content as the PDF, plus the format option in the API."""
from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from .conftest import generate_questions
from .test_papers_and_export_api import _paper, _token_of

pytestmark = pytest.mark.asyncio

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _all_text(docx_bytes: bytes) -> str:
    doc = Document(BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


async def _export_docx(client, paper_id, **body):
    response = await client.post(f"/export/{paper_id}", json={"format": "docx", **body})
    assert response.status_code == 200, response.text
    meta = response.json()
    download = await client.get(
        f"/export/files/{meta['filename']}", params={"token": _token_of(meta["download_url"])}
    )
    assert download.status_code == 200
    return meta, download


async def test_export_as_word_returns_a_real_docx(client):
    paper, _ = await _paper(client, count=2, marks=3)
    meta, download = await _export_docx(client, paper["id"])

    assert meta["filename"].endswith(".docx")
    assert download.headers["content-type"] == DOCX_MIME
    assert "attachment" in download.headers["content-disposition"]
    assert download.content.startswith(b"PK")  # a docx is a zip
    text = _all_text(download.content)
    assert "Karnataka State Board" in text
    assert "Answer Key" in text


async def test_word_answer_key_offers_any_n_of_m_points(client):
    paper, _ = await _paper(client, count=1, marks=3)
    _, download = await _export_docx(client, paper["id"])
    assert "any 3 of the following 6 points" in _all_text(download.content)


async def test_word_export_can_omit_the_answer_key(client):
    paper, _ = await _paper(client, count=2, marks=3)
    _, download = await _export_docx(client, paper["id"], include_answer_key=False)
    assert "Answer Key" not in _all_text(download.content)


async def test_default_format_is_still_pdf(client):
    paper, _ = await _paper(client, count=1, marks=3)
    response = await client.post(f"/export/{paper['id']}", json={})
    assert response.json()["filename"].endswith(".pdf")


async def test_unknown_export_format_is_rejected(client):
    paper, _ = await _paper(client, count=1, marks=3)
    response = await client.post(f"/export/{paper['id']}", json={"format": "rtf"})
    assert response.status_code == 422 or response.status_code == 400


@pytest.mark.parametrize("qtype,marks", [("MCQ", 1), ("Fill", 1), ("Match", 3), ("Long", 5), ("Short", 2)])
async def test_word_export_handles_every_question_type(client, qtype, marks):
    paper, _ = await _paper(client, count=2, marks=marks, type=qtype)
    _, download = await _export_docx(client, paper["id"])
    doc = Document(BytesIO(download.content))
    assert len(doc.paragraphs) > 5
